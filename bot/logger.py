import os
from datetime import datetime
import gspread
from google.oauth2.service_account import Credentials
from bot import config
import asyncio
import io
from openpyxl import Workbook
import json

_client = None

def get_gspread_client(force_refresh=False):
    """Возвращает кэшированный gspread клиент, избегая повторной авторизации на каждый запрос"""
    global _client
    if _client is not None and not force_refresh:
        return _client
    if not os.path.exists(config.CREDENTIALS_FILE):
        return None
    try:
        scopes = ['https://www.googleapis.com/auth/spreadsheets', 'https://www.googleapis.com/auth/drive']
        creds = Credentials.from_service_account_file(config.CREDENTIALS_FILE, scopes=scopes)
        _client = gspread.authorize(creds)
        return _client
    except Exception as e:
        print(f"Ошибка авторизации gspread: {e}")
        return None

def get_spreadsheet_id(guild_id=None):
    """Получение ID таблицы из настроек или конфига"""
    if os.path.exists('settings.json'):
        try:
            with open('settings.json', 'r') as f:
                settings = json.load(f)
                if guild_id:
                    g_str = str(guild_id)
                    if "guilds" in settings and g_str in settings["guilds"]:
                        if 'spreadsheet_id' in settings["guilds"][g_str]:
                            return settings["guilds"][g_str]['spreadsheet_id']
                if 'spreadsheet_id' in settings:
                    return settings['spreadsheet_id']
        except: pass
    return config.SPREADSHEET_ID


def get_sheet_by_type(client, table_type, guild_id=None):
    """Открывает и возвращает gspread.Worksheet для указанного типа таблицы с учетом настроек"""
    settings = {}
    if os.path.exists('settings.json'):
        try:
            with open('settings.json', 'r') as f:
                settings = json.load(f)
        except: pass
        
    sheet_config = {}
    if guild_id:
        g_str = str(guild_id)
        if "guilds" in settings and g_str in settings["guilds"]:
            sheet_config = settings["guilds"][g_str].get('sheet_config', {}).get(table_type, {})
            
    if not sheet_config:
        sheet_config = settings.get('sheet_config', {}).get(table_type, {})
    
    spreadsheet_id = sheet_config.get('spreadsheet_id')
    if not spreadsheet_id:
        if guild_id:
            g_str = str(guild_id)
            if "guilds" in settings and g_str in settings["guilds"]:
                spreadsheet_id = settings["guilds"][g_str].get('spreadsheet_id')
        if not spreadsheet_id:
            spreadsheet_id = settings.get('spreadsheet_id', config.SPREADSHEET_ID)
        
    spreadsheet = client.open_by_key(spreadsheet_id)
    
    gid = sheet_config.get('gid')
    if gid is not None:
        try:
            return spreadsheet.get_worksheet_by_id(int(gid))
        except Exception as e:
            print(f"Не удалось открыть лист {table_type} по gid {gid}: {e}. Пробуем по имени/дефолту...")
            
    # Дефолтные настройки, если gid не задан или не найден
    if table_type == 'call':
        return spreadsheet.sheet1
    elif table_type == 'zvz':
        try:
            return spreadsheet.get_worksheet_by_id(2085645540)
        except:
            try:
                return spreadsheet.worksheet("zvz")
            except:
                return spreadsheet.sheet1
    elif table_type == 'roum':
        try:
            return spreadsheet.get_worksheet_by_id(439097534)
        except:
            try:
                return spreadsheet.worksheet("roum")
            except:
                return spreadsheet.sheet1
    elif table_type == 'uf':
        try:
            return spreadsheet.worksheet("UF")
        except gspread.exceptions.WorksheetNotFound:
            return spreadsheet.add_worksheet(title="UF", rows="1000", cols="2")
    elif table_type == 'balance':
        return spreadsheet.worksheet("Balance")
    elif table_type == 'LogBalance':
        try:
            return spreadsheet.worksheet("LogBalance")
        except gspread.exceptions.WorksheetNotFound:
            return spreadsheet.add_worksheet(title="LogBalance", rows="1000", cols="6")
    elif table_type == 'LogSplit':
        try:
            return spreadsheet.worksheet("LogSplit")
        except gspread.exceptions.WorksheetNotFound:
            return spreadsheet.add_worksheet(title="LogSplit", rows="1000", cols="8")
    elif table_type == 'grev':
        try:
            return spreadsheet.get_worksheet_by_id(791255482)
        except Exception:
            try:
                return spreadsheet.worksheet("Grev")
            except gspread.exceptions.WorksheetNotFound:
                return spreadsheet.add_worksheet(title="Grev", rows="1000", cols="6")
    else:
        raise ValueError(f"Неизвестный тип таблицы: {table_type}")


_pending_saves = {}
_save_tasks = {}

async def save_data_async(call_data):
    """Асинхронное сохранение в Google Таблицы с дебаунсом (2 сек),
    чтобы объединять частые клики участников в один пакетный запрос"""
    call_id = call_data.get('id')
    if not call_id:
        return await asyncio.to_thread(save_to_google_sheets, call_data)
        
    _pending_saves[call_id] = call_data
    
    if call_id in _save_tasks and not _save_tasks[call_id].done():
        return
        
    async def _debounced_worker(c_id):
        try:
            await asyncio.sleep(2)
            data = _pending_saves.pop(c_id, None)
            if data:
                await asyncio.to_thread(save_to_google_sheets, data)
        except Exception as err:
            print(f"Ошибка дебаунса сохранения сбора {c_id}: {err}")
        finally:
            _save_tasks.pop(c_id, None)
            
    task = asyncio.create_task(_debounced_worker(call_id))
    _save_tasks[call_id] = task

async def save_members_to_uf_async(members, guild_id=None):
    """Асинхронное сохранение всех участников в таблицу UF"""
    return await asyncio.to_thread(save_members_to_uf, members, guild_id=guild_id)

def save_members_to_uf(members, guild_id=None):
    """Добавление только новых участников в Google Таблицу 'UF' и поиск тех, кто покинул сервер"""
    if not os.path.exists(config.CREDENTIALS_FILE):
        return [], 0, []
    
    try:
        client = get_gspread_client()
        if not client:
            return [], 0, []
        
        sheet = get_sheet_by_type(client, "uf", guild_id=guild_id)

        existing_data = sheet.get_all_values()
        existing_map = {row[0]: row[1] for row in existing_data if row and row[0].isdigit()}
        
        discord_ids = {str(m.id) for m in members}
        
        new_members = []
        rows_to_append = []
        for m in members:
            if str(m.id) not in existing_map:
                new_members.append(m)
                rows_to_append.append([str(m.id), m.display_name])
        
        if rows_to_append:
            sheet.append_rows(rows_to_append)
            
        missing_from_discord = []
        for s_id, s_nick in existing_map.items():
            if s_id not in discord_ids:
                missing_from_discord.append({"id": s_id, "nick": s_nick})
                
        total_count = len(existing_data) + len(rows_to_append)
        return new_members, total_count, missing_from_discord
            
    except Exception as e:
        print(f"Ошибка сохранения участников в UF: {e}")
        raise e

async def add_content_shares_async(text, admin_id=None, admin_name=None, guild_id=None):
    """Асинхронное добавление долей за контент"""
    return await asyncio.to_thread(add_content_shares, text, admin_id, admin_name, guild_id=guild_id)

def add_content_shares(text, admin_id=None, admin_name=None, guild_id=None):
    """Добавление долей в таблицу Balance и лог в LogSplit"""
    if not os.path.exists(config.CREDENTIALS_FILE):
        return [], []
    
    try:
        client = get_gspread_client()
        if not client:
            return [], []
        
        sheet = get_sheet_by_type(client, "balance", guild_id=guild_id)

        sheet.insert_cols([[""] * sheet.row_count], col=5)
        
        lines = [line.strip() for line in text.split("\n") if line.strip()]
        if len(lines) < 2:
            return [], []

        updates = []
        missing_users = []
        success_users = []
        log_rows = []
        
        updates.append({'range': 'E1', 'values': [[lines[0]]]})
        updates.append({'range': 'E2', 'values': [[lines[1]]]})
        
        ids_in_col_a = sheet.col_values(1)
        
        date_str = datetime.now().strftime("%Y-%m-%d %H:%M")
        split_name = lines[0]
        split_organizer = lines[1]
        
        for line in lines[2:]:
            parts = line.split()
            if len(parts) < 2: continue
            
            user_id = parts[0]
            amount_str = parts[-1]
            
            try:
                clean_amount = amount_str.replace(' ', '').replace(',', '.')
                if '.' in clean_amount:
                    amount = float(clean_amount)
                else:
                    amount = int(clean_amount)
            except ValueError:
                amount = amount_str

            try:
                row_idx = ids_in_col_a.index(user_id) + 1
                updates.append({'range': f'E{row_idx}', 'values': [[amount]]})
                nick_info = " ".join(parts[1:-1])
                if not nick_info:
                    nick_info = "Участник"
                success_users.append(f"Ник: {nick_info} ID: {user_id} Сумма: {amount}")
                
                # Сохраняем информацию для лога в LogSplit
                log_rows.append([
                    date_str,
                    str(admin_name) if admin_name else "Система",
                    str(admin_id) if admin_id else "0",
                    str(nick_info),
                    str(user_id),
                    str(amount),
                    str(split_name),
                    str(split_organizer)
                ])
            except ValueError:
                nick_info = " ".join(parts[1:-1])
                missing_users.append(f"Ник: {nick_info} ID: {user_id}")
        
        if updates:
            sheet.batch_update(updates, value_input_option='USER_ENTERED')
            try:
                sheet.format("E1:E", {
                    "textFormat": {
                        "fontSize": 14,
                        "bold": True
                    }
                })
            except Exception as e:
                print(f"Ошибка форматирования: {e}")
            
        # Запись логов сплитов в LogSplit
        if log_rows:
            try:
                log_sheet = get_sheet_by_type(client, "LogSplit", guild_id=guild_id)
                
                all_log_values = log_sheet.get_all_values()
                if not all_log_values or len(all_log_values) == 0:
                    log_sheet.append_row(["Дата", "Записал Ник", "Записал ID", "Кому Ник", "Кому ID", "Сумма", "Название сплита", "Организатор"])
                
                # Вставляем новые логи в начало (row=1), сдвигая прошлые вниз
                log_sheet.insert_rows(log_rows, row=1, value_input_option='USER_ENTERED')
            except Exception as e:
                print(f"Ошибка записи логов сплитов в LogSplit: {e}")
            
        return success_users, missing_users

    except Exception as e:
        print(f"Ошибка при добавлении долей: {e}")
        raise e

LOCAL_EXCEL_FILE = "calls_log.xlsx"

def save_to_local_excel(call_data):
    """Резервное сохранение данных сбора в локальный Excel-файл calls_log.xlsx"""
    try:
        from openpyxl import load_workbook
        if os.path.exists(LOCAL_EXCEL_FILE):
            wb = load_workbook(LOCAL_EXCEL_FILE)
        else:
            wb = Workbook()
            if "Sheet" in wb.sheetnames:
                wb.remove(wb["Sheet"])

        c_type = call_data.get('call_type', 'call')
        sheet_title = str(c_type)[:30]
        if sheet_title in wb.sheetnames:
            ws = wb[sheet_title]
        else:
            ws = wb.create_sheet(title=sheet_title)

        target_col = None
        max_col = ws.max_column
        for c in range(1, max_col + 1, 3):
            val = ws.cell(row=1, column=c).value
            if val is not None and str(val) == str(call_data['id']):
                target_col = c
                break

        if target_col is None:
            if ws.cell(row=1, column=1).value is None:
                target_col = 1
            else:
                ws.insert_cols(1, 3)
                target_col = 1

        date_str = datetime.now().strftime("%Y-%m-%d %H:%M")
        rows = [
            [str(call_data['id']), date_str, ""],
            [str(call_data['organizer'].id), str(call_data['organizer'].display_name), "Организатор"]
        ]
        for slot in call_data.get('slots', []):
            if isinstance(slot, list):
                for user in slot:
                    if user:
                        rows.append([str(user.id), user.display_name, ""])
            elif slot:
                rows.append([str(slot.id), slot.display_name, ""])

        # Очищаем блок перед записью
        for r in range(1, 60):
            for offset in range(3):
                ws.cell(row=r, column=target_col + offset, value=None)

        for r_idx, row_vals in enumerate(rows, 1):
            for c_idx, val in enumerate(row_vals):
                ws.cell(row=r_idx, column=target_col + c_idx, value=val)

        wb.save(LOCAL_EXCEL_FILE)
        wb.close()
    except Exception as err:
        print(f"Ошибка сохранения в локальный Excel ({LOCAL_EXCEL_FILE}): {err}")

def clear_local_excel_data(call_data):
    """Очистка сбора из локального Excel при удалении"""
    if not os.path.exists(LOCAL_EXCEL_FILE):
        return
    try:
        from openpyxl import load_workbook
        wb = load_workbook(LOCAL_EXCEL_FILE)
        c_type = call_data.get('call_type', 'call')
        sheet_title = str(c_type)[:30]
        if sheet_title in wb.sheetnames:
            ws = wb[sheet_title]
            max_col = ws.max_column
            target_col = None
            for c in range(1, max_col + 1, 3):
                val = ws.cell(row=1, column=c).value
                if val is not None and str(val) == str(call_data['id']):
                    target_col = c
                    break
            if target_col is not None:
                ws.delete_cols(target_col, 3)
                wb.save(LOCAL_EXCEL_FILE)
        wb.close()
    except Exception as err:
        print(f"Ошибка очистки локального Excel: {err}")

def save_to_google_sheets(call_data):
    """Сохранение данных сбора в локальный Excel (всегда) и в Google Таблицу (при наличии доступа)"""
    # 1. Резервная запись в локальный Excel, чтобы данные сбора никогда не потерялись
    save_to_local_excel(call_data)

    if not os.path.exists(config.CREDENTIALS_FILE):
        return
    
    try:
        client = get_gspread_client()
        if not client:
            return
            
        sheet = get_sheet_by_type(client, call_data.get('call_type', 'call'), guild_id=call_data.get('guild_id'))
        
        # Читаем ТОЛЬКО 1-ю строку для поиска ID сбора вместо выкачивания десятков тысяч ячеек get_all_values()!
        try:
            first_row = sheet.row_values(1)
        except Exception:
            first_row = []
            
        target_col_idx = -1
        if first_row:
            for col_i in range(0, len(first_row), 3):
                if first_row[col_i] == str(call_data['id']):
                    target_col_idx = col_i
                    break
        
        # Подсчитаем общее число участников
        total_users = 0
        for slot in call_data['slots']:
            if isinstance(slot, list):
                total_users += len(slot)
            elif slot:
                total_users += 1

        if target_col_idx == -1:
            if not first_row or not first_row[0]:
                target_col_idx = 0
            else:
                sheet.insert_cols([[""] * (total_users + 20)] * 3, col=1)
                target_col_idx = 0
            
        date_str = datetime.now().strftime("%Y-%m-%d %H:%M")

        rows = [
            [str(call_data['id']), date_str, ""],
            [str(call_data['organizer'].id), str(call_data['organizer'].display_name), "Организатор"]
        ]
        
        for slot in call_data['slots']:
            if isinstance(slot, list):
                for user in slot:
                    if user:
                        rows.append([str(user.id), user.display_name, ""])
            elif slot:
                rows.append([str(slot.id), slot.display_name, ""])
        
        # Дополняем хвост пустыми строками, чтобы затереть старых выписавшихся участников
        rows_to_write = rows + [["", "", ""]] * 15
        
        start_a1 = gspread.utils.rowcol_to_a1(1, target_col_idx + 1)
        end_a1 = gspread.utils.rowcol_to_a1(len(rows_to_write), target_col_idx + 3)
        sheet.update(range_name=f"{start_a1}:{end_a1}", values=rows_to_write)
            
    except Exception as e:
        print(f"Ошибка сохранения в Google Таблицы: {e}")

async def clear_data_async(call_data):
    """Асинхронная очистка данных при удалении сбора"""
    await asyncio.to_thread(clear_google_sheets_data, call_data)

def clear_google_sheets_data(call_data):
    clear_local_excel_data(call_data)
    if not os.path.exists(config.CREDENTIALS_FILE): return
    try:
        client = get_gspread_client()
        if not client: return
        sheet = get_sheet_by_type(client, call_data.get('call_type', 'call'), guild_id=call_data.get('guild_id'))
        
        try:
            first_row = sheet.row_values(1)
        except Exception:
            return
            
        target_col_idx = -1
        for col_i in range(0, len(first_row), 3):
            if first_row[col_i] == str(call_data['id']):
                target_col_idx = col_i
                break
        
        if target_col_idx != -1:
            start_a1 = gspread.utils.rowcol_to_a1(1, target_col_idx + 1)
            end_a1 = gspread.utils.rowcol_to_a1(50, target_col_idx + 3)
            empty_block = [["", "", ""]] * 50
            sheet.update(range_name=f"{start_a1}:{end_a1}", values=empty_block)
    except Exception as e:
        print(f"Ошибка при очистке Google Таблиц: {e}")

async def get_balance_uf_async(guild_id=None):
    """Получение всех данных (A, B, C) из таблицы Balance"""
    return await asyncio.to_thread(get_balance_uf, guild_id=guild_id)

def get_balance_uf(guild_id=None):
    if not os.path.exists(config.CREDENTIALS_FILE):
        return []
    try:
        client = get_gspread_client()
        if not client: return []
        sheet = get_sheet_by_type(client, "balance", guild_id=guild_id)
        
        all_rows = sheet.get_all_values()
        result = [row[:3] for row in all_rows if len(row) >= 3]
        return result
    except Exception as e:
        print(f"Ошибка get_balance_uf: {e}")
        return []

async def get_user_balance_async(user_id, guild_id=None):
    """Получение баланса конкретного пользователя по ID"""
    return await asyncio.to_thread(get_user_balance, user_id, guild_id=guild_id)

def get_user_balance(user_id, guild_id=None):
    if not os.path.exists(config.CREDENTIALS_FILE):
        return None
    try:
        client = get_gspread_client()
        if not client: return None
        sheet = get_sheet_by_type(client, "balance", guild_id=guild_id)
        
        ids_col = sheet.col_values(1)
        try:
            row_idx = ids_col.index(str(user_id)) + 1
            return sheet.cell(row_idx, 3).value
        except ValueError:
            return None
    except Exception as e:
        print(f"Ошибка get_user_balance: {e}")
        return None

async def export_balance_sheet_async(guild_id=None):
    """Экспорт текущего состояния таблицы Balance в Excel файл для бэкапа"""
    return await asyncio.to_thread(export_balance_sheet, guild_id=guild_id)

def export_balance_sheet(guild_id=None):
    if not os.path.exists(config.CREDENTIALS_FILE):
        return None
    try:
        client = get_gspread_client()
        if not client: return None, None
        sheet = get_sheet_by_type(client, "balance", guild_id=guild_id)
        
        data = sheet.get_all_values()
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Balance Backup"
        
        for r_idx, row in enumerate(data, 1):
            for c_idx, value in enumerate(row, 1):
                ws.cell(row=r_idx, column=c_idx, value=value)
        
        filename = f"balance_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        wb.close()
        
        return buffer, filename
    except Exception as e:
        print(f"Ошибка export_balance_sheet: {e}")
        return None, None

async def update_balance_rows_async(members, guild_id=None):
    """Асинхронное обновление строк в таблице Balance"""
    return await asyncio.to_thread(update_balance_rows, members, guild_id=guild_id)

def update_balance_rows(members, guild_id=None):
    """Добавление новых участников в таблицу Balance с формулами"""
    if not os.path.exists(config.CREDENTIALS_FILE):
        return 0
    try:
        client = get_gspread_client()
        if not client: return 0
        sheet = get_sheet_by_type(client, "balance", guild_id=guild_id)
        
        existing_ids = set(sheet.col_values(1))
        
        rows_to_append = []
        new_count = 0
        
        current_rows = len(sheet.get_all_values())
        
        for m in members:
            if str(m.id) not in existing_ids:
                new_count += 1
                row_idx = current_rows + new_count
                formula = '=SUM(INDIRECT("D"&ROW()&":AAC"&ROW()))'
                rows_to_append.append([str(m.id), m.display_name, formula])
        
        if rows_to_append:
            sheet.append_rows(rows_to_append, value_input_option='USER_ENTERED')
            
            try:
                start_row = current_rows + 1
                end_row = current_rows + new_count
                sheet.format(f"A{start_row}:C{end_row}", {
                    "textFormat": {
                        "fontSize": 14,
                        "bold": True
                    }
                })
            except Exception as e:
                print(f"Ошибка форматирования новых строк: {e}")
            
        return new_count
    except Exception as e:
        print(f"Ошибка update_balance_rows: {e}")
        return 0

async def archive_balance_sheet_async(guild_id=None):
    """Асинхронная архивация балансов"""
    return await asyncio.to_thread(archive_balance_sheet, guild_id=guild_id)

def archive_balance_sheet(guild_id=None):
    """Перенос C -> D, очистка E+, очистка E1, E2"""
    if not os.path.exists(config.CREDENTIALS_FILE): return False
    try:
        client = get_gspread_client()
        if not client: return False
        sheet = get_sheet_by_type(client, "balance", guild_id=guild_id)
        
        all_values = sheet.get_all_values()
        if not all_values: return False
        
        c_values = sheet.col_values(3)
        d_updates = []
        for val in c_values:
            if not val:
                d_updates.append([""])
                continue
            
            clean_val = val.replace(' ', '').replace(',', '.')
            try:
                if '.' in clean_val:
                    num_val = float(clean_val)
                else:
                    num_val = int(clean_val)
                d_updates.append([num_val])
            except:
                d_updates.append([val])
        
        sheet.update(range_name=f"D1:D{len(c_values)}", values=d_updates, value_input_option='USER_ENTERED')
        
        max_cols = sheet.col_count
        if max_cols >= 5:
            range_to_clear = f"E1:{gspread.utils.rowcol_to_a1(len(all_values), max_cols)}"
            sheet.batch_clear([range_to_clear])
        
        num_rows = len(all_values)
        if num_rows >= 3:
            formulas = []
            for i in range(3, num_rows + 1):
                formulas.append(['=SUM(INDIRECT("D"&ROW()&":AAC"&ROW()))'])
            sheet.update(range_name=f"C3:C{num_rows}", values=formulas, value_input_option='USER_ENTERED')

        sort_end_cell = gspread.utils.rowcol_to_a1(num_rows, max_cols)
        sheet.sort((2, 'asc'), range=f'A3:{sort_end_cell}')

        return True
    except Exception as e:
        print(f"Ошибка archive_balance_sheet: {e}")
        return False

async def withdraw_user_balance_async(user_id, treasurer_id, treasurer_name, guild_id=None):
    """Асинхронное снятие денег (обнуление) и запись лога в LogBalance"""
    return await asyncio.to_thread(withdraw_user_balance, user_id, treasurer_id, treasurer_name, guild_id=guild_id)

def withdraw_user_balance(user_id, treasurer_id, treasurer_name, guild_id=None):
    """Обнуление баланса пользователя: C->Old, D=0, Clear E+, и добавление лога в LogBalance"""
    if not os.path.exists(config.CREDENTIALS_FILE): return None
    try:
        client = get_gspread_client()
        if not client: return None
        sheet = get_sheet_by_type(client, "balance", guild_id=guild_id)
        
        ids_col = sheet.col_values(1)
        try:
            row_idx = ids_col.index(str(user_id)) + 1
            target_nick = sheet.cell(row_idx, 2).value or "Неизвестно"
            old_balance = sheet.cell(row_idx, 3).value
            
            max_cols = sheet.col_count
            if max_cols >= 4:
                range_to_clear = f"{gspread.utils.rowcol_to_a1(row_idx, 4)}:{gspread.utils.rowcol_to_a1(row_idx, max_cols)}"
                sheet.batch_clear([range_to_clear])
            
            sheet.update_cell(row_idx, 4, 0)
            
            # Запись лога снятия в LogBalance
            log_sheet = get_sheet_by_type(client, "LogBalance", guild_id=guild_id)
                
            # Проверим, пустой ли лист, чтобы записать заголовки
            all_log_values = log_sheet.get_all_values()
            if not all_log_values or len(all_log_values) == 0:
                log_sheet.append_row(["Дата", "Казначей Ник", "Казначей ID", "Чей счет Ник", "Чей счет ID", "Сумма"])
                
            date_str = datetime.now().strftime("%Y-%m-%d %H:%M")
            log_sheet.insert_row([
                date_str,
                str(treasurer_name),
                str(treasurer_id),
                str(target_nick),
                str(user_id),
                str(old_balance)
            ], index=1, value_input_option='USER_ENTERED')
            
            return old_balance
        except ValueError:
            return None
    except Exception as e:
        print(f"Ошибка withdraw_user_balance: {e}")
        return None

async def sort_uf_sheet_async(guild_id=None):
    """Асинхронная сортировка UF"""
    return await asyncio.to_thread(sort_uf_sheet, guild_id=guild_id)

def sort_uf_sheet(guild_id=None):
    """Сортировка листа UF по нику (Столбец B)"""
    if not os.path.exists(config.CREDENTIALS_FILE): return False
    try:
        client = get_gspread_client()
        if not client: return False
        sheet = get_sheet_by_type(client, "uf", guild_id=guild_id)
        
        all_values = sheet.get_all_values()
        if len(all_values) <= 1: return True
        
        sheet.sort((2, 'asc'), range='A2:B' + str(len(all_values)))
        return True
    except Exception as e:
        print(f"Ошибка sort_uf_sheet: {e}")
        return False

async def get_attendance_stats_async(days=7, guild_id=None):
    """Асинхронная статистика посещений по всем 3 типам контента"""
    return await asyncio.to_thread(get_attendance_stats, days, guild_id)

def get_attendance_stats(days=7, guild_id=None):
    """Подсчет посещений за последние N дней по типам сборов"""
    if not os.path.exists(config.CREDENTIALS_FILE): return {}
    try:
        client = get_gspread_client()
        if not client: return {}
        
        from datetime import datetime, timedelta
        now = datetime.now()
        threshold = now - timedelta(days=days)
        
        stats = {} # {id: {"nick": nick, "call": 0, "zvz": 0, "roum": 0, "organizer": 0}}
        
        # Получаем все 3 листа
        types = [
            ('call', get_sheet_by_type(client, 'call', guild_id)), 
            ('zvz', get_sheet_by_type(client, 'zvz', guild_id)), 
            ('roum', get_sheet_by_type(client, 'roum', guild_id))
        ]
        
        for c_type, sheet in types:
            try:
                all_values = sheet.get_all_values()
            except Exception as e:
                print(f"Ошибка чтения листа {c_type}: {e}")
                continue
                
            if not all_values or len(all_values[0]) == 0: continue
            
            first_row = all_values[0]
            for col_i in range(0, len(first_row), 3):
                if col_i + 1 >= len(first_row): continue
                date_str = first_row[col_i + 1]
                if not date_str: continue
                
                try:
                    call_date = datetime.strptime(date_str, "%Y-%m-%d %H:%M")
                    if call_date < threshold: continue
                    
                    organizer_id = None
                    if len(all_values) > 1 and len(all_values[1]) > col_i + 1:
                        organizer_id = all_values[1][col_i]
                        organizer_nick = all_values[1][col_i + 1]
                        if organizer_id and organizer_id.isdigit():
                            if organizer_id not in stats:
                                stats[organizer_id] = {"nick": organizer_nick, "call": 0, "zvz": 0, "roum": 0, "organizer": 0}
                            stats[organizer_id]["organizer"] += 1
                        
                    for r_idx, row in enumerate(all_values[2:], start=2):
                        if len(row) > col_i + 1:
                            u_id = row[col_i]
                            u_nick = row[col_i + 1]
                            
                            if u_id and u_id.isdigit():
                                if u_id not in stats:
                                    stats[u_id] = {"nick": u_nick, "call": 0, "zvz": 0, "roum": 0, "organizer": 0}
                                stats[u_id][c_type] += 1
                except Exception as e:
                    continue
                    
        return stats
    except Exception as e:
        print(f"Ошибка get_attendance_stats: {e}")
        return {}


def sync_grev_data(txt_content=None, guild_id=None):
    """Синхронизация данных Grev:
    1. Загрузка локального состояния state_{guild_id}.json
    2. Загрузка данных из Google таблицы Grev
    3. Синхронизация Discord ID и ручных изменений баланса
    4. Применение логов из txt_content (если переданы)
    5. Запись результатов обратно в таблицу и сохранение state_{guild_id}.json
    """
    import json
    import os
    import re
    
    # 1. Загрузка локального состояния
    STATE_FILE = f"state_{guild_id}.json" if guild_id else "state.json"
    state = {
        "last_transaction": None,
        "balances": {},
        "discord_ids": {},
        "player_last_tx": {}
    }
    
    if os.path.exists(STATE_FILE):
        try:
            with open(STATE_FILE, 'r', encoding='utf-8') as f:
                loaded = json.load(f)
                if isinstance(loaded, dict):
                    state["last_transaction"] = loaded.get("last_transaction")
                    state["balances"] = loaded.get("balances", {})
                    state["discord_ids"] = loaded.get("discord_ids", {})
                    state["player_last_tx"] = loaded.get("player_last_tx", {})
        except Exception as e:
            print(f"[Предупреждение] Ошибка при чтении {STATE_FILE}: {e}")

    if not os.path.exists(config.CREDENTIALS_FILE):
        raise FileNotFoundError(f"Файл credentials.json не найден.")

    # 2. Подключение к Google Sheets
    client = get_gspread_client()
    if not client:
        raise ConnectionError("Не удалось подключиться к Google Таблицам")
    sheet = get_sheet_by_type(client, "grev", guild_id)
    
    # 3. Чтение данных из Google Sheets
    all_values = sheet.get_all_values()
    sheet_players = {} # nickname -> info
    
    if all_values:
        for row in all_values[1:]:
            if len(row) < 2 or not row[1].strip():
                continue
            discord_id = row[0].strip()
            nickname = row[1].strip()
            balance_str = row[2].strip() if len(row) > 2 else ""
            date_tx = row[3].strip() if len(row) > 3 else ""
            action_tx = row[4].strip() if len(row) > 4 else ""
            amount_tx = row[5].strip() if len(row) > 5 else ""
            
            try:
                balance_str_clean = balance_str.replace(' ', '').replace(',', '.')
                balance = int(float(balance_str_clean)) if balance_str_clean else 0
            except ValueError:
                balance = 0
                
            sheet_players[nickname] = {
                "discord_id": discord_id,
                "balance": balance,
                "last_tx": {
                    "date": date_tx,
                    "action": action_tx,
                    "amount": amount_tx
                }
            }

    # 4. Синхронизация ручных правок из Google Sheets в локальный state.json
    # a) Discord IDs (Google Sheet имеет приоритет)
    for nick, p_data in sheet_players.items():
        if p_data["discord_id"]:
            state["discord_ids"][nick] = p_data["discord_id"]
            
    # b) Балансы (ручные изменения в таблице перезаписывают локальный стейт)
    for nick, p_data in sheet_players.items():
        if nick not in state["balances"]:
            state["balances"][nick] = p_data["balance"]
        elif p_data["balance"] != state["balances"][nick]:
            state["balances"][nick] = p_data["balance"]
            
        # Синхронизируем инфо о последней операции
        if nick not in state["player_last_tx"]:
            state["player_last_tx"][nick] = p_data["last_tx"]

    # 5. Обработка файла логов, если предоставлен
    new_tx_count = 0
    if txt_content:
        lines = [line.strip() for line in txt_content.split('\n') if line.strip()]
        all_transactions = []
        start_idx = 1 if lines and ("Дата" in lines[0] or "Игрок" in lines[0] or "Date" in lines[0] or "Player" in lines[0]) else 0
        
        for line in lines[start_idx:]:
            parts = line.split('\t')
            if len(parts) < 4:
                parts = re.findall(r'"([^"]*)"', line)
                if len(parts) < 4:
                    continue
            clean_parts = [p.strip('"') for p in parts]
            
            date = clean_parts[0]
            player = clean_parts[1]
            action = clean_parts[2]
            
            try:
                amount = int(clean_parts[3])
            except ValueError:
                try:
                    amount = float(clean_parts[3])
                except ValueError:
                    amount = 0
                    
            all_transactions.append({
                "date": date,
                "player": player,
                "action": action,
                "amount": amount
            })
            
        if all_transactions:
            last_tx = state.get("last_transaction")
            new_transactions = []
            if last_tx:
                found_idx = -1
                for idx, tx in enumerate(all_transactions):
                    if (tx["date"] == last_tx.get("date") and 
                        tx["player"] == last_tx.get("player") and 
                        tx["action"] == last_tx.get("action") and 
                        tx["amount"] == last_tx.get("amount")):
                        found_idx = idx
                        break
                if found_idx != -1:
                    new_transactions = all_transactions[:found_idx]
                else:
                    new_transactions = all_transactions
            else:
                new_transactions = all_transactions
                
            if new_transactions:
                new_tx_count = len(new_transactions)
                for tx in reversed(new_transactions):
                    player = tx["player"]
                    amount = tx["amount"]
                    state["balances"][player] = state["balances"].get(player, 0) + amount
                    
                    state["player_last_tx"][player] = {
                        "date": tx["date"],
                        "action": tx["action"],
                        "amount": tx["amount"]
                    }
                    
                state["last_transaction"] = new_transactions[0]

    # 6. Сохраняем стейт локально
    try:
        with open(STATE_FILE, 'w', encoding='utf-8') as f:
            json.dump(state, f, ensure_ascii=False, indent=4)
    except Exception as e:
        print(f"[Ошибка] Не удалось сохранить state.json: {e}")
        
    # 7. Записываем обновленные балансы обратно в Google Sheets
    headers = ['ID Discord', 'Игровой Ник', 'Баланс', 'Дата последней операции', 'Действие последней операции', 'Сумма последней операции']
    rows_to_write = [headers]
    
    sorted_players = sorted(state["balances"].items(), key=lambda x: x[1], reverse=True)
    
    for nick, balance in sorted_players:
        discord_id = state["discord_ids"].get(nick, "")
        player_tx = state["player_last_tx"].get(nick, {})
        
        rows_to_write.append([
            str(discord_id),
            str(nick),
            int(balance),
            str(player_tx.get("date", "")),
            str(player_tx.get("action", "")),
            str(player_tx.get("amount", ""))
        ])
        
    sheet.clear()
    sheet.update(range_name=f"A1:F{len(rows_to_write)}", values=rows_to_write, value_input_option='USER_ENTERED')
    
    try:
        sheet.format("A1:F1", {
            "textFormat": {
                "bold": True,
                "fontSize": 11
            }
        })
    except Exception as e:
        print(f"Ошибка форматирования: {e}")
        
    return new_tx_count


async def sync_grev_data_async(txt_content=None, guild_id=None):
    """Асинхронная синхронизация балансов Грев"""
    import asyncio
    return await asyncio.to_thread(sync_grev_data, txt_content, guild_id)
